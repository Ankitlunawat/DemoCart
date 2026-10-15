import unittest,csv,json
from pathlib import Path
from openpyxl import load_workbook
from run import run
class PipelineTests(unittest.TestCase):
 def test_offline_exports_and_quality(self):
  rows,report=run();out=Path(__file__).parent/'output'
  self.assertEqual(len(rows),60);self.assertEqual(report['unique_skus'],60)
  self.assertEqual(json.loads((out/'products.json').read_text()),rows)
  with (out/'products.csv').open(encoding='utf-8-sig') as f:self.assertEqual(len(list(csv.DictReader(f))),60)
  wb=load_workbook(out/'products.xlsx');self.assertEqual(wb.active.max_row,61);self.assertEqual(wb.active.freeze_panes,'A2')

if __name__=="__main__":unittest.main()
